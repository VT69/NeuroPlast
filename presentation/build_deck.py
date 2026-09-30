"""Final-review slide deck (16:9 PowerPoint), built only from the frozen results.

    python presentation/build_deck.py        # -> presentation/NeuroPlast_final_review.pptx

Every number comes from the paper's numbers pipeline (paper/compute_numbers.py, via demo/build_dashboard.py), the
charts are drawn from the same data as the dashboard, and the figures are the paper's (paper/figures/*.png, 300 dpi).
Palette = the dashboard's method colours. Speaker notes are on every slide.
"""
from __future__ import annotations

import os
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "demo"))
import build_dashboard as bd  # noqa: E402  (chdirs to ROOT; exposes the paper's numbers and the dashboard data)

OUT = "presentation/NeuroPlast_final_review.pptx"
FIG = "paper/figures"

# ---------------------------------------------------------------- palette (dashboard tokens)
def rgb(h):
    return RGBColor.from_string(h)


INK, INK2, MUTED, HAIR = "151514", "52514E", "85837D", "E6E5E0"
WHITE, CARD, CARD2, DARK = "FFFFFF", "F4F4F1", "EBEAE4", "151514"
SLEEP, REPLAY, ISO, EWC, NAIVE, MATCHED = "2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "4A3AA7"
BADGE = {"helps": ("E3F4EC", "0F6B4A", "◐"), "tradeoff": ("FBF0D9", "7A5300", "⇄"),
         "none": ("EFEEE9", "52514E", "○"), "untested": ("EFEEE9", "52514E", "?")}
HEAD, BODY = "Cambria", "Calibri"
W, H = 13.333, 7.5
M = 0.6  # side margin (in)

N = bd.numbers()
mean = lambda k: N[k].split(" ±")[0]  # noqa: E731


# ---------------------------------------------------------------- helpers
def txt(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, margin=0.0):
    """paras: list of dicts {t, size, bold, color, font, italic, space_after, bullet} or (runs list under 'runs')."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        runs = p.get("runs") or [dict(t=p["t"])]
        for r in runs:
            run = para.add_run()
            run.text = r["t"]
            f = run.font
            f.name = r.get("font", p.get("font", BODY))
            f.size = Pt(r.get("size", p.get("size", 16)))
            f.bold = r.get("bold", p.get("bold", False))
            f.italic = r.get("italic", p.get("italic", False))
            f.color.rgb = rgb(r.get("color", p.get("color", INK)))
        if p.get("space_after") is not None:
            para.space_after = Pt(p["space_after"])
        if p.get("line") is not None:
            para.line_spacing = p["line"]
        if p.get("bullet"):
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.28))))
            pPr.set("indent", str(-Emu(Inches(0.22))))
            for tag in ("a:buNone", "a:buChar", "a:buFont", "a:buClr"):
                for e in pPr.findall(qn(tag)):
                    pPr.remove(e)
            clr = pPr.makeelement(qn("a:buClr"), {})
            srgb = clr.makeelement(qn("a:srgbClr"), {"val": p.get("bullet_color", MUTED)})
            clr.append(srgb)
            pPr.append(clr)
            pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
    return tb


def bullets(slide, x, y, w, h, items, size=17, color=INK, gap=10):
    paras = []
    for it in items:
        if isinstance(it, str):
            paras.append(dict(t=it, size=size, color=color, bullet=True, space_after=gap))
        else:  # list of runs
            paras.append(dict(runs=it, size=size, color=color, bullet=True, space_after=gap))
    return txt(slide, x, y, w, h, paras)


def card(slide, x, y, w, h, fill=CARD, radius=0.06, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def pill(slide, x, y, text, kind, size=12, w=None):
    bg, fg, icon = BADGE[kind]
    w = w or (0.45 + 0.085 * len(text) * size / 12)
    s = card(slide, x, y, w, 0.34 * size / 12 + 0.06, fill=bg, radius=0.5)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = f"{icon}  {text}"
    r.font.size, r.font.bold, r.font.name = Pt(size), True, BODY
    r.font.color.rgb = rgb(fg)
    return s


def image(slide, path, x, y, w, h):
    from PIL import Image
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    dw, dh = iw * s, ih * s
    return slide.shapes.add_picture(path, Inches(x + (w - dw) / 2), Inches(y + (h - dh) / 2), Inches(dw), Inches(dh))


def title(slide, text, kicker=None, dark=False):
    y = 0.45
    if kicker:
        txt(slide, M, y, W - 2 * M, 0.35, [dict(t=kicker.upper(), size=12, bold=True, color=SLEEP if not dark else "9CC3F2")])
        y += 0.36
    txt(slide, M, y, W - 2 * M, 0.9, [dict(t=text, size=30, bold=True, font=HEAD, color=WHITE if dark else INK)])


def footer(slide, n, total, dark=False):
    c = "8F8D86" if dark else MUTED
    txt(slide, M, H - 0.45, 6, 0.3, [dict(t="NeuroPlast · final review", size=10, color=c)])
    txt(slide, W - M - 2, H - 0.45, 2, 0.3, [dict(t=f"{n} / {total}", size=10, color=c, align=PP_ALIGN.RIGHT)])


def stat(slide, x, y, w, big, label, color=INK, big_size=40):
    txt(slide, x, y, w, 0.75, [dict(t=big, size=big_size, bold=True, font=HEAD, color=color)])
    txt(slide, x, y + 0.72, w, 0.8, [dict(t=label, size=13, color=INK2, line=1.05)])


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = rgb(color)


def arrow(slide, x1, y1, x2, y2, color="A9A79E", width=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def box(slide, x, y, w, h, head, lines, fill=CARD, accent=None, head_size=16, line_size=12.5):
    card(slide, x, y, w, h, fill=fill, radius=0.12)
    if accent:
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.22), Inches(y + 0.24), Inches(0.14), Inches(0.14))
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb(accent)
        dot.line.fill.background()
    hx = x + (0.44 if accent else 0.22)
    txt(slide, hx, y + 0.14, w - (hx - x) - 0.15, 0.4, [dict(t=head, size=head_size, bold=True)])
    txt(slide, x + 0.22, y + 0.55, w - 0.4, h - 0.6, [dict(t=l, size=line_size, color=INK2, space_after=2) for l in lines])


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = " ".join(text.split())


# ---------------------------------------------------------------- slides
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    slides = []

    def new(dark=False):
        s = prs.slides.add_slide(blank)
        bg(s, DARK if dark else WHITE)
        slides.append((s, dark))
        return s

    fg = bd.forgetting()
    mem = bd.memory()
    key = bd.key_memory(N, mem)

    # 1 ---- title
    s = new(dark=True)
    txt(s, M + 0.1, 1.0, 11, 0.4, [dict(t="CAPSTONE PROJECT · FINAL REVIEW", size=13, bold=True, color="9CC3F2")])
    txt(s, M + 0.1, 1.5, 11.6, 2.0, [dict(t="Do Brain-Inspired Mechanisms Help Continual Reinforcement Learning?",
                                           size=40, bold=True, font=HEAD, color=WHITE, line=1.0)])
    txt(s, M + 0.1, 3.45, 11, 0.5, [dict(t="A controlled ablation study with spiking agents", size=22, color="C3C2B7")])
    rows = [("Vaibhav Tiwari", True), ("Roll number: [to fill in]", False), ("Guide: [to fill in]", False),
            ("Department: [to fill in]", False), ("Institution: [to fill in]", False), ("Date: [to fill in]", False)]
    txt(s, M + 0.1, 4.55, 7, 2.4, [dict(t=t, size=18 if b else 15, bold=b, color=WHITE if b else "C3C2B7", space_after=3)
                                   for t, b in rows])
    notes(s, """Good morning. My project asks a simple question: when people add brain-inspired mechanisms to a
    reinforcement-learning agent, such as spiking neurons, local learning rules or a sleep phase, which of them actually
    help? I built one agent, switched each mechanism on and off against a matched control, and measured the effect. The
    short answer, which I will justify over the next fourteen slides, is: mostly no, with two specific exceptions.
    Every number in this talk comes from the frozen results and is the same number that appears in the paper.""")

    # 2 ---- problem
    s = new()
    title(s, "Learning a new task erases the old one", "The problem")
    bullets(s, M, 1.9, 5.6, 4.5, [
        "An agent learns tasks one after another: fetch the red ball, then the green key, then the blue box",
        "Plain fine-tuning overwrites what earlier tasks needed: catastrophic forgetting",
        "Standard fixes: penalise weight changes (EWC), replay stored experience, or one network per task",
        "Brain-inspired fixes (spikes, local plasticity, sleep) are popular, but are usually added all at once",
    ], size=17)
    cnn = fg["cnn3"]
    cd = CategoryChartData()
    cd.categories = [f"after task {i}" for i in range(1, cnn["T"] + 1)]
    cols = {}
    for srs in cnn["series"]:
        if srs["key"] in ("naive", "replay", "sleep"):
            cd.add_series(srs["label"], [p["mean"] for p in srs["first"]])
            cols[srs["label"]] = {"naive": NAIVE, "replay": REPLAY, "sleep": SLEEP}[srs["key"]]
    card(s, 6.75, 1.85, 6.0, 4.75, fill="FAFAF8", radius=0.05)
    txt(s, 7.0, 2.0, 5.6, 0.4, [dict(t="Return on task 1 as later tasks are learned", size=15, bold=True)])
    txt(s, 7.0, 2.35, 5.6, 0.35, [dict(t=f"CNN agent, 3 tasks, mean of {cnn['series'][0]['n']} seeds; replay (dashed) and "
                                         "sleep overlap near 1.0", size=12, color=MUTED)])
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(6.9), Inches(2.7), Inches(5.7), Inches(3.8), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size, ch.legend.font.name = Pt(12), BODY
    ch.legend.font.color.rgb = rgb(INK2)
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 1.0, 0.25
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(HAIR)
    va.format.line.fill.background()
    va.tick_labels.font.size, va.tick_labels.font.name = Pt(11), BODY
    va.tick_labels.font.color.rgb = rgb(MUTED)
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.00", False
    ca = ch.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(11), BODY
    ca.tick_labels.font.color.rgb = rgb(MUTED)
    ca.format.line.color.rgb = rgb(HAIR)
    for srs in ch.plots[0].series:
        c = cols[srs.name]
        srs.smooth = False
        srs.format.line.color.rgb = rgb(c)
        srs.format.line.width = Pt(2.75)
        if srs.name == "replay":
            srs.format.line.dash_style = MSO_LINE.DASH
        srs.marker.style = XL_MARKER_STYLE.CIRCLE
        srs.marker.size = 8
        srs.marker.format.fill.solid()
        srs.marker.format.fill.fore_color.rgb = rgb(c)
        srs.marker.format.line.color.rgb = rgb(WHITE)
    notes(s, f"""Continual learning means the agent learns tasks in sequence. In my setting the agent sees a small room
    with three objects and must fetch a different object in each task. With plain fine-tuning, learning the second and
    third tasks destroys the first: on the right, the naive agent's return on task 1 drops from about
    {cnn['series'][0]['first'][0]['mean']:.2f} to {cnn['series'][0]['first'][-1]['mean']:.2f}. Replay and sleep keep it
    near 1. The standard remedies are regularisation like EWC, replaying stored experience, or giving each task its own
    network. Brain-inspired remedies are popular, but papers usually add several at once, so it is hard to say which
    one did the work. That is the gap this project addresses.""")

    # 3 ---- research question
    s = new()
    title(s, "Which brain-inspired mechanisms actually help?", "Research question")
    txt(s, M, 1.95, W - 2 * M, 0.8, [dict(t="Switch each mechanism on alone and compare it with a control that is not "
                                           "brain-inspired but matches its resources.", size=18, color=INK2)])
    mechs = [("Spiking encoder", "same-shape CNN, matched operations"),
             ("STDP (local plasticity)", "a random update of the same size"),
             ("Homeostasis", "the same agent without it"),
             ("DFA: learning without backprop", "ordinary backpropagation"),
             ("Sleep-like consolidation", "replay with an exactly matched replay budget"),
             ("One shared network", "one network per task, at equal total memory"),
             ("Transformer working memory", "no memory, and a stack of recent frames")]
    cw, chh, gx, gy = 2.87, 1.5, 0.25, 0.22
    for i, (m, c) in enumerate(mechs):
        r, k = divmod(i, 4)
        x = M + k * (cw + gx) + (0 if r == 0 else (cw + gx) / 2)
        y = 2.8 + r * (chh + gy)
        card(s, x, y, cw, chh, fill=CARD, radius=0.1)
        txt(s, x + 0.2, y + 0.14, cw - 0.4, 0.62, [dict(t=m, size=15, bold=True, line=0.95)])
        txt(s, x + 0.2, y + 0.8, cw - 0.4, 0.66, [dict(runs=[dict(t="vs  ", bold=True, color=SLEEP), dict(t=c)], size=13, color=INK2)])
    txt(s, M, 6.35, W - 2 * M, 0.5, [dict(t="Negative results count: the aim is to find out which mechanisms contribute, "
                                           "not to make all of them win.", size=14, italic=True, color=INK2)])
    notes(s, """The research question is: which of these mechanisms actually helps, when each one is isolated? Each card
    shows a mechanism and the control it is compared against. The controls are chosen to remove the most obvious
    alternative explanation. For example, if STDP helps, is it because STDP is useful, or because any extra update of
    that size helps? So STDP is compared with a random update of exactly the same size. If sleep helps, is it because
    of sleep, or because it replays more old data? So sleep is compared with replay at the same number of replayed
    samples. The project was set up so that negative results are results, not failures.""")

    # 4 ---- related work
    s = new()
    title(s, "Related work", "Where this sits")
    rw = [("Continual RL", ["EWC: protect important weights (Kirkpatrick et al., 2017)",
                            "CLEAR: replay + behavioural cloning, a very strong baseline (Rolnick et al., 2019)",
                            "PackNet: isolate parameters per task (Mallya & Lazebnik, 2018)"], ISO),
          ("Sleep and consolidation", ["Sleep-like unsupervised replay recovers forgotten tasks in classifiers "
                                       "(Tadros et al., 2022)",
                                       "Ours: offline replay + distillation in RL, compared with CLEAR-style replay"], SLEEP),
          ("Spiking networks for RL", ["Surrogate gradients make SNNs trainable (Neftci et al., 2019)",
                                       "The surrogate slope matters more in RL (Van den Berghe et al., 2025)",
                                       "Spiking decision transformers (Pandey & Biswas, 2025)"], MATCHED),
          ("Local learning rules", ["Three-factor, reward-modulated STDP (Frémaux & Gerstner, 2016)",
                                    "Direct feedback alignment: no weight transport (Nøkland, 2016)",
                                    "Homeostatic plasticity (Turrigiano & Nelson, 2004)"], REPLAY)]
    cw, chh = (W - 2 * M - 0.3) / 2, 2.0
    for i, (hd, ls, col) in enumerate(rw):
        r, k = divmod(i, 2)
        box(s, M + k * (cw + 0.3), 1.9 + r * (chh + 0.25), cw, chh, hd, ls, accent=col, line_size=13.5)
    txt(s, M, 6.4, W - 2 * M, 0.5, [dict(runs=[dict(t="Gap: ", bold=True), dict(
        t="these mechanisms are usually tested together, against an unmodified baseline, so it is unclear which one did the work.")],
        size=14, color=INK2)])
    notes(s, """Four strands of work. In continual RL, the standard baselines are EWC, replay, and parameter isolation;
    the CLEAR paper showed that replay with behavioural cloning is extremely strong, which is why it is my main
    comparator for sleep. Tadros and colleagues showed a sleep-like phase can recover forgotten tasks in classifiers; my
    sleep phase is inspired by that but is replay plus distillation in RL. On the spiking side, surrogate gradients make
    SNNs trainable, and recent work shows RL is sensitive to how they are set. Local rules like three-factor STDP, DFA and
    homeostasis are proposed as brain-like alternatives to backprop. The gap: they are usually tested together.""")

    # 5 ---- architecture
    s = new()
    title(s, "System architecture", "What I built")
    A = dict(T=N["snnT"], window=N["hybridWindow"], buffer=N["bufferPerTask"], period=N["sleepPeriod"], steps=N["sleepSteps"])
    bw, bh, by = 2.75, 1.55, 1.95
    xs = [M + i * (bw + 0.42) for i in range(4)]
    comps = [("Environment", ["MiniGrid, 7×7 view of the room", "task ID given each episode", "reward for the right object"], None),
             ("Spiking encoder", ["3 conv layers of LIF neurons", f"{A['T']} spike steps per frame", "control: same-shape CNN"], SLEEP),
             ("Transformer memory", ["1 layer, 4 attention heads", f"attends over the last {A['window']} frames",
                                     "SNN+Transformer agent only"], MATCHED),
             ("Actor-critic heads", ["policy + value head per task", "task embedding on features", "trained with PPO"], None)]
    for x, (hd, ls, col) in zip(xs, comps):
        box(s, x, by, bw, bh, hd, ls, accent=col, line_size=12.5)
    for i in range(3):
        arrow(s, xs[i] + bw + 0.05, by + bh / 2, xs[i + 1] - 0.05, by + bh / 2)
    ly = by + bh + 0.35
    for seg in ((xs[3] + bw / 2, by + bh + 0.05, xs[3] + bw / 2, ly),):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *[Inches(v) for v in seg])
        c.line.color.rgb, c.line.width = rgb("A9A79E"), Pt(1.75)
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xs[3] + bw / 2), Inches(ly), Inches(xs[0] + bw / 2), Inches(ly))
    c.line.color.rgb, c.line.width = rgb("A9A79E"), Pt(1.75)
    arrow(s, xs[0] + bw / 2, ly, xs[0] + bw / 2, by + bh + 0.05)
    lab = card(s, W / 2 - 0.5, ly - 0.17, 1.0, 0.34, fill=WHITE, radius=0.5)
    txt(s, W / 2 - 0.5, ly - 0.17, 1.0, 0.34, [dict(t="action", size=12, color=MUTED, align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, M, 4.45, 12, 0.35, [dict(t="Continual training of the sleep variant: wake and sleep phases alternate on the same shared network",
                                    size=13, color=MUTED)])
    py, ph = 4.9, 1.6
    box(s, M, py, 4.6, ph, "Wake: online learning", ["PPO on the current task, acting in the environment",
                                                     "states stored in the buffer at the end of each task"], accent=INK2)
    cyl = s.shapes.add_shape(MSO_SHAPE.CAN, Inches(5.75), Inches(py + 0.1), Inches(1.85), Inches(ph - 0.2))
    cyl.fill.solid()
    cyl.fill.fore_color.rgb = rgb(CARD2)
    cyl.line.fill.background()
    txt(s, 5.75, py + 0.55, 1.85, 0.9, [dict(t="Replay buffer", size=14, bold=True, align=PP_ALIGN.CENTER),
                                        dict(t=f"{A['buffer']} states / task", size=12, color=INK2, align=PP_ALIGN.CENTER)])
    box(s, 8.13, py, W - M - 8.13, ph, "Sleep: offline consolidation",
        [f"every {A['period']} updates and at each task end: {A['steps']} gradient steps, no environment",
         "distil old tasks' policies + rehearse the current task"], accent=SLEEP)
    arrow(s, M + 4.6 + 0.05, py + ph / 2, 5.7, py + ph / 2)
    arrow(s, 7.65, py + ph / 2, 8.08, py + ph / 2)
    notes(s, f"""This is the agent. The environment gives a 7 by 7 view of the room and the task ID. The encoder is a
    spiking network: three convolutional layers of leaky integrate-and-fire neurons, simulated for {A['T']} time steps
    per frame; the control is an ordinary CNN with the same shape. Optionally a small Transformer attends over the last
    {A['window']} frames. Actor-critic heads, one per task, are trained with PPO. For continual learning, the sleep
    variant alternates two phases: in the wake phase the agent learns the current task online; at the end of each task
    it stores {A['buffer']} states in a replay buffer. Every {A['period']} updates it sleeps: {A['steps']} offline
    gradient steps that distil its old policies from the buffer, with no environment interaction. Every part can be
    swapped independently, which is what makes the ablations possible.""")

    # 6 ---- setup + fairness
    s = new()
    title(s, "Experimental setup and fairness controls", "How it was tested")
    box(s, M, 1.9, 5.7, 3.3, "Setup", [], accent=INK2)
    bullets(s, M + 0.22, 2.5, 5.3, 2.7, [
        "MiniGrid: fetch3 / fetch5 sequences (task ID given), DoorKey-6x6, Memory S7/S11/S13",
        f"PPO, small agents ({N['paramsCnnPerTask']}–{N['paramsHybridShared']} parameters)",
        "Metrics: final accuracy (ACC), forgetting, learning speed (AUC), solve rate",
    ], size=15, gap=8)
    for k, (big, lab, col) in enumerate([(N["seedRange"], "seeds per arm", INK), (N["computeRuns"], "PPO runs", INK),
                                          (N["computeCoreHours"], "single-thread CPU hours, no GPU", INK)]):
        stat(s, M + 0.1 + k * 1.95, 5.4, 1.85, big, lab, color=col, big_size=34)
    box(s, 6.6, 1.9, W - M - 6.6, 4.65, "Fairness controls", [], accent=SLEEP)
    bullets(s, 6.82, 2.5, W - M - 6.9, 4.0, [
        "Same frames, seeds and learning rate in every arm of a comparison",
        "STDP vs a random update of exactly the same size",
        "Sleep vs replay with the number of replayed samples matched exactly",
        "Memory counted as parameters plus replay buffer, not parameters alone",
        "Pre-registered tests; Welch / Fisher; Holm correction; 95% CIs",
    ], size=16, gap=12)
    notes(s, f"""The experiments use MiniGrid. The continual sequences are fetch3 and fetch5, where the agent is told
    which task it is doing. DoorKey is used for the plasticity questions, and the Memory maps for the Transformer
    question. All agents are small and train with PPO on a CPU; in total the frozen results are {N['computeRuns']} runs
    and about {N['computeCoreHours']} CPU hours. The right-hand side is the part I want to stress. Each comparison keeps
    everything else equal, and in two places I matched a resource that is usually ignored: the number of replayed
    samples, and total memory including the replay buffer. The final tests were written down before the runs, with
    Holm correction where several tests share a family.""")

    # 7 ---- scorecard
    s = new()
    title(s, "Results at a glance: mostly no, with two exceptions", "Mechanism scorecard")
    short = {
        "Spiking encoder": f"works at {N['rqThreeSnnTwoLowOps']} ops/frame where the CNN needs {N['rqThreeCnnFloorOps']}; "
                           f"{N['snnCpuSlowdown']}× CPU cost",
        "Local learning rule (STDP)": f"vs random update: ΔAUC {N['rqOneContrBAucCi']}, p = {N['rqOneContrBAucP']}",
        "Homeostasis, without backprop": f"ΔAUC {N['dfaAucCi']}, Holm p = {N['dfaAucHolm']}; solved "
                                         f"{N['rqOneDfaHomeoSolved']} vs {N['rqOneDfaSolved']} (n.s.)",
        "Homeostasis, with backprop": f"ΔAUC {N['rqOneContrCAucCi']}; Holm p = {N['rqOneContrCAucHolm']}",
        "Sleep-like consolidation": f"sleep − replay ΔACC {N['matchAccCi']}, p = {N['matchAccP']}",
        "One shared network": f"200 states/task: ties isolation in {N['smallSnnTwoHundredMb']} vs {N['memSnnThreeIsoMb']} MB",
        "Transformer working memory": "no learnable memory benchmark within budget",
    }
    y0, rh = 1.9, 0.66
    txt(s, M + 0.2, y0 - 0.02, 3.6, 0.3, [dict(t="MECHANISM", size=10.5, bold=True, color=MUTED)])
    txt(s, 4.2, y0 - 0.02, 3.4, 0.3, [dict(t="VERDICT", size=10.5, bold=True, color=MUTED)])
    txt(s, 7.85, y0 - 0.02, 5, 0.3, [dict(t="KEY NUMBER (95% CI)", size=10.5, bold=True, color=MUTED)])
    for i, r in enumerate(bd.scorecard(N)):
        y = y0 + 0.35 + i * rh
        card(s, M, y, W - 2 * M, rh - 0.1, fill=CARD if i % 2 == 0 else "FAFAF8", radius=0.2)
        txt(s, M + 0.2, y, 3.5, rh - 0.1, [dict(t=r["mech"], size=14.5, bold=True)], anchor=MSO_ANCHOR.MIDDLE)
        pill(s, 4.2, y + 0.1, r["verdict"], r["kind"], size=12, w=3.45)
        txt(s, 7.85, y, W - M - 7.95, rh - 0.1, [dict(t=short[r["mech"]], size=13, color=INK2)], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, """This is the whole result on one slide. Three mechanisms had no effect beyond their control: STDP, which
    was indistinguishable from a random update of the same size; homeostasis with ordinary backprop, which looked
    positive but did not survive the multiple-comparison correction; and sleep, which removes forgetting but ties replay
    once the replay budget is matched. The spiking encoder is a trade-off: it works at fewer operations, but costs more
    CPU time to train and the energy savings are only an estimate. The two exceptions are homeostasis when the network
    is trained without backprop, and one shared network when the replay buffer is small. The Transformer question could
    not be tested. The next slides go through each row.""")

    # 8 ---- forgetting
    s = new()
    title(s, "Sleep and replay stop forgetting; sleep does not beat replay", "Forgetting")
    image(s, f"{FIG}/fig_forgetting.png", M - 0.1, 1.8, 8.3, 4.9)
    x0 = 9.0
    stat(s, x0, 1.9, 3.8, mean("accCnnNaive"), "naive fine-tuning, final ACC (CNN, 3 tasks)", color=NAIVE)
    stat(s, x0, 3.35, 3.8, f"{mean('accCnnReplay')} · {mean('accCnnSleep')}", "replay · sleep, final ACC", color=SLEEP, big_size=34)
    stat(s, x0, 4.8, 3.8, N["matchAccCi"].split(" [")[0], f"sleep − replay at an exactly matched replay budget, "
                                                          f"95% CI {N['matchAccCi'].split(' ', 1)[1]}, p = {N['matchAccP']}",
         big_size=34)
    notes(s, f"""The top row is the return on the first task as later tasks are learned; the bottom row is the average
    over tasks seen so far. Naive fine-tuning, in pink, forgets: its final accuracy on the 3-task CNN sequence is
    {mean('accCnnNaive')}. Replay and sleep both remove almost all of that forgetting, reaching {mean('accCnnReplay')}
    and {mean('accCnnSleep')}. On the SNN plus Transformer agent sleep originally looked slightly better than replay,
    but it was replaying {N['unmatchedExtraPct']} percent more data. In the pre-registered test with the replayed samples
    matched exactly, the difference is {N['matchAccCi']}, p = {N['matchAccP']}, with {N['matchN']} seeds per arm. So in
    my implementation sleep is a way of scheduling replay, and at equal replay it is not better.""")

    # 9 ---- memory
    s = new()
    title(s, "A shared network wins only with a small replay buffer", "Memory trade-off")
    image(s, f"{FIG}/fig_memory.png", M - 0.1, 1.85, 8.5, 3.9)
    txt(s, M, 5.85, 8.3, 0.9, [dict(t="Total memory = fp32 parameters + replay buffer. Numbers above the axis: replay states "
                                       "stored per task.", size=12, color=MUTED)])
    x0 = 9.2
    stat(s, x0, 1.9, 3.6, f"{key['ratioPct']}%", f"of the memory: SNN sleep with {key['buf']} states/task ties isolation "
                                                 f"({key['sleepAcc']} vs {key['isoAcc']}; {key['sleepMb']} vs {key['isoMb']} MB)",
         color=SLEEP, big_size=48)
    stat(s, x0, 3.55, 3.6, f"{N['perParamRatioRange']}×", "more accuracy per parameter for the shared network", big_size=34)
    stat(s, x0, 4.95, 3.6, f"{N['perMbRatioRange']}×", f"more accuracy per MB for isolation at the default "
                                                      f"{N['bufferPerTask']} states/task", color=ISO, big_size=34)
    notes(s, f"""Isolation, one network per task, never forgets by construction, so the fair question is memory. Per
    parameter, the shared network looks {N['perParamRatioRange']} times more efficient, which is the usual way this is
    reported. But the shared network needs a replay buffer, and at the default {N['bufferPerTask']} states per task the
    buffer outweighs the extra networks: isolation gets {N['perMbRatioRange']} times more accuracy per megabyte. The
    exception is a small buffer: on the SNN, sleep with {key['buf']} states per task ties isolation's accuracy,
    {key['sleepAcc']} against {key['isoAcc']}, in {key['sleepMb']} instead of {key['isoMb']} megabytes, about
    {key['ratioPct']} percent of the memory. With three seeds the confidence interval is wide, so this is 'no large
    difference', not proof of equality.""")

    # 10 ---- local plasticity
    s = new()
    title(s, "STDP adds nothing; homeostasis speeds up DFA", "Local plasticity")
    image(s, f"{FIG}/fig_rq1.png", M - 0.1, 1.85, 7.9, 3.9)
    txt(s, M, 5.85, 7.6, 0.8, [dict(t=f"DoorKey-6x6, {N['rqOneBpN']} seeds per arm, pre-registered. Dots: seeds (filled = "
                                       f"solved); bars: mean and 95% CI; numbers: seeds solved.", size=12, color=MUTED)])
    bullets(s, 8.55, 1.95, W - M - 8.55, 4.6, [
        [dict(t="STDP ≈ random update: ", bold=True), dict(t=f"ΔAUC {N['rqOneContrBAucCi']}, p = {N['rqOneContrBAucP']}")],
        [dict(t="Unstabilised STDP hurts: ", bold=True), dict(t=f"solved {N['vanillaStdpSolved']} vs {N['vanillaBpSolved']} "
                                                              "with backprop (runaway firing)")],
        [dict(t="Homeostasis + backprop: ", bold=True), dict(t=f"raw p = {N['rqOneContrCAucP']}, Holm p = "
                                                             f"{N['rqOneContrCAucHolm']}: not significant")],
        [dict(t="Homeostasis + DFA: ", bold=True), dict(t=f"ΔAUC {N['dfaAucCi']}, Holm p = {N['dfaAucHolm']}; solved "
                                                        f"{N['rqOneDfaHomeoSolved']} vs {N['rqOneDfaSolved']} "
                                                        f"(Holm p = {N['dfaFisherHolm']})")],
    ], size=15, gap=12)
    notes(s, f"""This is DoorKey, a task where the encoder's learning matters. On the left, backprop is the learning
    signal. Adding stabilised, reward-modulated STDP does not beat a random update of the same size: the AUC difference
    is {N['rqOneContrBAucCi']}, p = {N['rqOneContrBAucP']}. Without stabilisation, STDP makes things worse through
    runaway firing. Homeostasis with backprop looked positive, raw p = {N['rqOneContrCAucP']}, but not after Holm
    correction. On the right the encoder learns with direct feedback alignment instead of backprop. Here homeostasis
    clearly speeds up learning: AUC difference {N['dfaAucCi']}, Holm p = {N['dfaAucHolm']}. It solved
    {N['rqOneDfaHomeoSolved']} runs against {N['rqOneDfaSolved']}, but that difference is not significant, so it makes
    DFA faster, not reliable. This is the first of the two exceptions.""")

    # 11 ---- spiking encoder
    s = new()
    title(s, "Spiking encoder: fewer operations, not more accuracy", "Efficiency")
    image(s, f"{FIG}/fig_frontier.png", M - 0.1, 1.85, 8.3, 3.6)
    txt(s, M, 5.5, 8.1, 1.0, [dict(t=f"Behaviour cloning of a DoorKey-6x6 teacher, activity penalty swept; {N['rqThreeSeeds']} "
                                      "seeds per point. Operations = synaptic events (SNN) or non-zero multiply-accumulates "
                                      "(sparsified CNN).", size=12, color=MUTED)])
    x0 = 9.2
    stat(s, x0, 1.9, 3.6, N["rqThreeSnnTwoLowOps"], f"ops/frame: the SNN still works (accuracy {N['rqThreeSnnTwoLowAcc']})",
         color=MATCHED, big_size=40)
    stat(s, x0, 3.35, 3.6, N["rqThreeCnnFloorOps"], "ops/frame: below this the sparsified CNN stops working", big_size=40)
    stat(s, x0, 4.8, 3.6, f"{N['snnCpuSlowdown']}×", f"CPU time per training step ({N['msSnn']} vs {N['msCnn']} ms); "
                                                   "energy savings are only an estimate", color=REPLAY, big_size=40)
    notes(s, f"""To compare encoders without RL noise, both were trained to copy a DoorKey teacher, with a penalty that
    trades operations for accuracy. Where both work, they are equally accurate. The difference is at the low end: the
    sparsified CNN collapses below about {N['rqThreeCnnFloorOps']} operations per frame, while the SNN still works at
    {N['rqThreeSnnTwoLowOps']}, with accuracy {N['rqThreeSnnTwoLowAcc']}. So the SNN reaches a lower-operation regime,
    not higher accuracy. Two caveats: on the CPU I actually used, the SNN is {N['snnCpuSlowdown']} times slower to train,
    because it simulates several time steps; and the energy numbers are estimates that assume neuromorphic hardware,
    which I did not measure.""")

    # 12 ---- didn't survive
    s = new()
    title(s, "What didn't survive controls, and why", "Honest accounting")
    rows = [
        ("Sleep beats replay (SNN+Transformer)", f"p = {N['hybSleepReplayNThreeP']} with 3 seeds",
         f"tie at a matched replay budget: p = {N['matchAccP']}, {N['matchN']} seeds", "seeds; replay budget"),
        ("Shared weights far more efficient", f"{N['perParamRatioRange']}× per parameter",
         f"isolation {N['perMbRatioRange']}× better per MB", "memory accounting"),
        ("Sleep far beats isolation, 5 tasks (SNN)", f"gap {N['fiveSnnShortCi']}",
         f"gap {N['fiveSnnFairCi']} at a fair budget", "training budget"),
        ("Homeostasis helps backprop", f"p = {N['homeoEarlyP']} (3 vs 5 seeds)",
         f"Holm p = {N['rqOneContrCAucHolm']} with {N['rqOneBpN']} seeds", "seeds; multiple tests"),
        ("Homeostasis makes DFA work", f"solved {N['dfaEarlySolvedHomeo']} vs {N['dfaEarlySolvedDfa']}",
         f"{N['rqOneDfaHomeoSolved']} vs {N['rqOneDfaSolved']}, p = {N['dfaFisherP']} (AUC effect survives)", "seeds"),
        ("Transformer helps on a memory task", f"S7 pilot {N['sSevenSnnTf']} vs {N['sSevenSnnEight']}",
         f"no memory needed on S7 (memoryless CNN {N['sSevenCnnOne']})", "benchmark validity"),
    ]
    cols_x = [M, 4.3, 7.0, 10.75]
    cols_w = [3.6, 2.6, 3.65, W - M - 10.75]
    for j, hd in enumerate(["EARLIER FINDING", "EVIDENCE THEN", "AFTER THE CONTROL", "CAUSE"]):
        txt(s, cols_x[j] + 0.15, 1.88, cols_w[j], 0.3, [dict(t=hd, size=10.5, bold=True, color=MUTED)])
    rh = 0.7
    for i, r in enumerate(rows):
        y = 2.25 + i * rh
        card(s, M, y, W - 2 * M, rh - 0.1, fill=CARD if i % 2 == 0 else "FAFAF8", radius=0.2)
        for j, t in enumerate(r):
            txt(s, cols_x[j] + 0.15, y, cols_w[j] - 0.2, rh - 0.1,
                [dict(t=t, size=13 if j else 13.5, bold=(j == 0), color=INK if j in (0, 3) else INK2)], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, M, 6.55, W - 2 * M, 0.4, [dict(t="Each apparent win disappeared once the resource it quietly used more of was matched.",
                                           size=14, italic=True, color=INK2)])
    notes(s, f"""This slide is, I think, the most useful part of the project. Six findings looked positive early on,
    and each shrank or disappeared once the right control was in place. Sleep beating replay came from three seeds and
    from sleep replaying more data. Shared weights looked far more efficient only because the replay buffer was not
    counted as memory. Sleep beating isolation on five tasks came from a training budget too small for fresh networks.
    The two homeostasis results came from small seed counts, and the Transformer result came from a benchmark that did
    not need memory at all. The common pattern: each apparent win disappeared when the resource it quietly used more of,
    whether seeds, frames, memory or replayed samples, was matched.""")

    # 13 ---- live demo
    s = new()
    title(s, "Live demo", "See it yourself")
    card(s, M, 1.85, 7.6, 4.75, fill=CARD, radius=0.04)
    image(s, "presentation/assets/dash_agents.png", M + 0.15, 1.95, 7.3, 4.55)
    box(s, 8.55, 1.85, W - M - 8.55, 2.1, "1. Results dashboard", [
        "open demo/index.html (one file, works offline)",
        "scorecard, agents, forgetting and memory charts"], accent=SLEEP, line_size=13)
    box(s, 8.55, 4.15, W - M - 8.55, 2.45, "2. Run an agent on CPU", [
        "python demo/live_episode.py",
        "   --agent naive --task 0 --window",
        "then the same with --agent sleep",
        "prints each step and the real accuracy matrix"], accent=ISO, line_size=13)
    notes(s, f"""For the demo I will switch to the dashboard, which is a single offline HTML file built from the same
    frozen results. It shows the two CNN agents after learning all three tasks, in the same rooms, chosen before seeing
    the outcomes. The naive agent's first task-1 episode happens to succeed, but its average task-1 return over 100
    evaluation episodes is only {bd.demo_agents()['naive']['R'][-1][0]:.3f}, while the sleep agent stays at
    {bd.demo_agents()['sleep']['R'][-1][0]:.3f}. If time allows I will also run one live episode on the CPU with
    live_episode.py, first the naive agent, then the sleep agent, on task 0.""")

    # 14 ---- limitations + future
    s = new()
    title(s, "Limitations and future work", "What this does not show")
    box(s, M, 1.85, 5.9, 4.3, "Limitations", [], accent=REPLAY)
    bullets(s, M + 0.22, 2.45, 5.5, 4.1, [
        "Task ID is given (task-incremental), so this says nothing about task-agnostic learning",
        f"Small scale: MiniGrid, ≤{N['paramsHybridShared']} parameters, 3–5 tasks",
        f"Few seeds: many continual arms have {N['nCnnSleep']}, so CIs are wide",
        "Energy is an estimate from operation counts; no hardware was measured",
        "Transformer memory untested: no benchmark was learnable in budget",
    ], size=16, gap=12)
    box(s, 6.75, 1.85, W - M - 6.75, 4.3, "Future work", [], accent=SLEEP)
    bullets(s, 6.97, 2.45, W - M - 7.1, 4.1, [
        "A memory benchmark agents can learn (the cue seen on the way to the choice)",
        [dict(t="Compressed replay (idea, untested): ", bold=True), dict(
            t="the buffer size decides the shared-vs-isolation trade-off, so shrink it: quantised targets, "
              "fewer but better-chosen states, or generative replay")],
        "Homeostasis target sweep: is the DFA effect robust?",
        "More seeds for the operation-count frontier",
    ], size=16, gap=12)
    notes(s, f"""The limitations are real. The agent is told which task it is doing, the scale is small, many
    continual arms have only {N['nCnnSleep']} seeds, the energy numbers are estimates, and the Transformer question is
    untested because no memory benchmark was learnable within budget. For future work, the most direct next step is a
    memory benchmark that agents can actually learn. The idea I find most promising is compressed replay: since the size
    of the replay buffer is what decides whether a shared network beats one network per task, making the buffer cheaper
    moves that trade-off directly, for example by storing quantised targets, fewer but better-chosen states, or a small
    generative model. This is an idea, not a result.""")

    # 15 ---- conclusion
    s = new(dark=True)
    txt(s, M + 0.1, 0.9, 11, 0.4, [dict(t="CONCLUSION", size=13, bold=True, color="9CC3F2")])
    txt(s, M + 0.1, 1.4, 11.8, 1.0, [dict(t="Mostly no, with two specific exceptions.", size=40, bold=True, font=HEAD,
                                         color=WHITE)])
    items = [("No mechanism beat its matched control", "on final accuracy or forgetting: STDP ≈ random update, sleep ≈ replay"),
             ("Two exceptions", "homeostasis speeds up learning without backprop; a shared network is more memory-"
                                "efficient with a small replay buffer"),
             ("The lesson is methodological", "every apparent win vanished once the resource it quietly used more "
                                              "of (seeds, frames, memory, replayed samples) was matched")]
    for i, (hd, body) in enumerate(items):
        y = 2.75 + i * 1.2
        num = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(M + 0.1), Inches(y + 0.02), Inches(0.5), Inches(0.5))
        num.fill.solid()
        num.fill.fore_color.rgb = rgb("2A2A28")
        num.line.fill.background()
        txt(s, M + 0.1, y + 0.02, 0.5, 0.5, [dict(t=str(i + 1), size=16, bold=True, color="9CC3F2", align=PP_ALIGN.CENTER)],
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, M + 0.85, y - 0.03, 11, 0.4, [dict(t=hd, size=19, bold=True, color=WHITE)])
        txt(s, M + 0.85, y + 0.38, 11, 0.6, [dict(t=body, size=15, color="C3C2B7")])
    txt(s, M + 0.1, 6.35, 11, 0.5, [dict(t="Thank you. Questions?", size=20, bold=True, color=WHITE)])
    notes(s, """To conclude. None of the brain-inspired mechanisms beat its matched control on accuracy or forgetting.
    There are two specific exceptions: homeostasis speeds up learning when the network is trained without backprop, and
    a shared network is more memory-efficient than one network per task when the replay buffer is small. The broader
    lesson is methodological: in this project, every apparent win of a biological mechanism disappeared once I matched
    the resource it was quietly using more of. My suggestion for this field is to report and match those budgets.
    Thank you; I'm happy to take questions. There are backup slides with the full tables and references.""")

    # B1 ---- continual table
    s = new()
    title(s, "Backup: final accuracy and forgetting by method", "Backup")
    cols = ["", "CNN, 3 tasks", "SNN, 3 tasks", "SNN+Transformer, 3 tasks", "CNN, 5 tasks"]
    acc = [("naive", "Naive"), ("EWC (λ=100)", "Ewc"), ("replay", "Replay"), ("sleep", "Sleep"),
           ("sleep, matched", "SleepMatched"), ("isolation", "Iso")]
    seqs = ["Cnn", "Snn", "Hyb", "CnnFive"]

    def g(prefix, m, q):
        return N.get(f"{prefix}{q}{m}", "–")
    data = [["Final ACC (mean ± SD)", "", "", "", ""]]
    data += [[lab] + [g("acc", m, q) for q in seqs] for lab, m in acc]
    data += [["FORGET (mean)", "", "", "", ""]]
    data += [[lab] + [g("forget", m, q) for q in seqs] for lab, m in acc if m in ("Naive", "Replay", "Sleep", "SleepMatched")]
    tbl_shape = s.shapes.add_table(len(data) + 1, 5, Inches(M), Inches(1.9), Inches(W - 2 * M), Inches(4.5))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(3.0)
    for j in range(1, 5):
        tbl.columns[j].width = Inches((W - 2 * M - 3.0) / 4)
    for i, row in enumerate([cols] + data):
        for j, v in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            head = i == 0 or row[0] in ("Final ACC (mean ± SD)", "FORGET (mean)")
            cell.fill.fore_color.rgb = rgb(CARD2 if i == 0 else (CARD if head else WHITE))
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            tf = cell.text_frame
            tf.text = ""
            r = tf.paragraphs[0].add_run()
            r.text = v
            r.font.size, r.font.name = Pt(13), BODY
            r.font.bold = head
            r.font.color.rgb = rgb(INK if head else INK2)
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
    txt(s, M, 6.6, W - 2 * M, 0.4, [dict(t=f"Seeds: {N['nCnnSleep']} per arm, except SNN+Transformer (replay {N['nHybReplay']}, "
                                           f"sleep {N['nHybSleep']}, sleep matched {N['nHybSleepMatched']}). Paper Table 3.",
                                           size=12, color=MUTED)])
    notes(s, """Backup: the full continual results table from the paper, final accuracy with standard deviation over
    seeds, and forgetting. Sleep, matched is the pre-registered budget-matched arm, run only on the SNN plus Transformer
    agent, where the replay budgets differed.""")

    # B2 ---- RQ1 contrasts
    s = new()
    title(s, "Backup: pre-registered local-plasticity contrasts", "Backup")
    hdr = ["Contrast", "Δ solve rate [95% CI]", "Fisher p (Holm)", "Δ AUC [95% CI]", "Welch p (Holm)"]
    rows = []
    for k, lab in (("A", "STDP vs homeostasis alone"), ("B", "STDP vs matched random"), ("C", "homeostasis vs backprop"),
                   ("D", "random vs homeostasis")):
        rows.append([lab, N[f"rqOneContr{k}SolveCi"], f"{N[f'rqOneContr{k}FisherP']} ({N[f'rqOneContr{k}FisherHolm']})",
                     N[f"rqOneContr{k}AucCi"], f"{N[f'rqOneContr{k}AucP']} ({N[f'rqOneContr{k}AucHolm']})"])
    rows.append(["DFA + homeostasis vs DFA", N["dfaSolveCi"], f"{N['dfaFisherP']} ({N['dfaFisherHolm']})", N["dfaAucCi"],
                 f"{N['dfaAucP']} ({N['dfaAucHolm']})"])
    tbl = s.shapes.add_table(len(rows) + 1, 5, Inches(M), Inches(2.0), Inches(W - 2 * M), Inches(3.3)).table
    widths = [3.4, 2.4, 1.9, 2.4, W - 2 * M - 10.1]
    for j, wd in enumerate(widths):
        tbl.columns[j].width = Inches(wd)
    for i, row in enumerate([hdr] + rows):
        for j, v in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(CARD2 if i == 0 else (CARD if i == len(rows) else WHITE))
            tf = cell.text_frame
            tf.text = ""
            r = tf.paragraphs[0].add_run()
            r.text = v
            r.font.size, r.font.name, r.font.bold = Pt(13), BODY, i == 0
            r.font.color.rgb = rgb(INK if i == 0 else INK2)
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
    txt(s, M, 5.55, W - 2 * M, 0.8, [dict(t=f"DoorKey-6x6, {N['rqOneBpN']} seeds per arm. First four rows: one Holm family per "
                                           "statistic (backprop as the global signal). Last row: the DFA family (pre-registered "
                                           "separately). Newcombe CIs for solve rates, Welch for AUC.", size=12, color=MUTED)])
    notes(s, """Backup: the exact pre-registered contrasts. The only effect that survives correction is the AUC effect
    of homeostasis under DFA, in the last row.""")

    # B3 ---- protocol
    s = new()
    title(s, "Backup: protocol constants", "Backup")
    prot = [("PPO", f"{N['numEnvs']} envs × {N['numSteps']} steps, {N['updateEpochs']} epochs, minibatch {N['minibatch']}, "
                    f"clip {N['clipCoef']}, γ = {N['gammaVal']}, GAE λ = {N['gaeLambda']}, KL stop {N['targetKl']}"),
            ("Learning rate", f"{N['lrFlat']} (CNN, SNN); {N['lrHybrid']} (SNN+Transformer)"),
            ("Frames per task", f"{N['framesPerTaskFlat']} (CNN, SNN); {N['framesPerTaskHybrid']} (SNN+Transformer); "
                                f"{N['framesPerTaskFiveFair']} (5-task SNN, fair budget)"),
            ("SNN", f"{N['snnT']} time steps, learnable leak (init {N['lifBeta']}), surrogate slope {N['surrogateSlope']}, "
                    "membrane reset each frame"),
            ("Plasticity", f"α = {N['stdpAlpha']}; homeostasis target {N['homeoTarget']}, rate {N['homeoLr']}"),
            ("Sleep", f"every {N['sleepPeriod']} updates + task end; {N['sleepSteps']} steps per phase; "
                      f"{N['bufferPerTask']} states per task"),
            ("Evaluation", f"{N['evalEpisodesCl']} episodes per task (continual); {N['evalEpisodesSingle']} (single task)"),
            ("Memory accounting", f"{N['bytesPerState']} B/state (CNN, SNN); {N['bytesPerStateHybrid']} B/state (SNN+Transformer)")]
    for i, (k, v) in enumerate(prot):
        y = 1.95 + i * 0.58
        card(s, M, y, W - 2 * M, 0.5, fill=CARD if i % 2 == 0 else "FAFAF8", radius=0.2)
        txt(s, M + 0.2, y, 2.6, 0.5, [dict(t=k, size=14, bold=True)], anchor=MSO_ANCHOR.MIDDLE)
        txt(s, 3.3, y, W - M - 3.4, 0.5, [dict(t=v, size=13.5, color=INK2)], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Backup: the protocol constants, read from the stored run configurations (paper appendix A).")

    # B4 ---- references
    s = new()
    title(s, "Backup: references", "Backup")
    refs = ["Bi & Poo (1998). Synaptic modifications in cultured hippocampal neurons. J. Neurosci.",
            "Chaudhry et al. (2018). Riemannian walk for incremental learning. ECCV.",
            "Chevalier-Boisvert et al. (2023). Minigrid & Miniworld. NeurIPS Datasets and Benchmarks.",
            "Frémaux & Gerstner (2016). Neuromodulated STDP and three-factor learning rules. Front. Neural Circuits.",
            "Holm (1979). A simple sequentially rejective multiple test procedure. Scand. J. Statistics.",
            "Khetarpal et al. (2022). Towards continual reinforcement learning: a review. JAIR.",
            "Kirkpatrick et al. (2017). Overcoming catastrophic forgetting in neural networks. PNAS.",
            "Mallya & Lazebnik (2018). PackNet. CVPR.",
            "Neftci, Mostafa & Zenke (2019). Surrogate gradient learning in SNNs. IEEE Signal Proc. Mag.",
            "Nøkland (2016). Direct feedback alignment. NeurIPS.",
            "Pandey & Biswas (2025). Spiking decision transformers. arXiv:2508.21505.",
            "Rolnick et al. (2019). Experience replay for continual learning (CLEAR). NeurIPS.",
            "Schulman et al. (2017). Proximal policy optimization algorithms. arXiv:1707.06347.",
            "Tadros et al. (2022). Sleep-like unsupervised replay reduces catastrophic forgetting. Nat. Commun.",
            "Turrigiano & Nelson (2004). Homeostatic plasticity in the developing nervous system. Nat. Rev. Neurosci.",
            "Van den Berghe et al. (2025). Adaptive surrogate gradients for sequential RL in SNNs. NeurIPS.",
            "Wołczyk et al. (2021). Continual World. NeurIPS."]
    half = (len(refs) + 1) // 2
    for k, chunk in enumerate((refs[:half], refs[half:])):
        txt(s, M + k * 6.1, 1.9, 5.9, 5.0, [dict(t=r, size=11.5, color=INK2, space_after=5) for r in chunk])
    notes(s, "Backup: the main references. The full list, with the verification status of each entry, is in "
             "paper/refs.bib.")

    total_main = 15
    for i, (sl, dark) in enumerate(slides):
        n = i + 1
        footer(sl, n if n <= total_main else f"B{n - total_main}", total_main, dark)
    prs.save(OUT)
    print("wrote", OUT, len(slides), "slides")


if __name__ == "__main__":
    build()
